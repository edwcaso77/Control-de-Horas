import datetime
import io

from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from funcionarios.models import Area
from funcionarios.permisos import administrador_requerido

from .clasificacion import FIN_NOCHE, INICIO_NOCHE
from .conteo import COLUMNAS_HORAS, CLAVES, conteo_semana, semana_de


def _semana_elegida(request):
    try:
        fecha = datetime.date.fromisoformat(request.GET.get('semana', ''))
    except ValueError:
        fecha = timezone.localdate()
    return semana_de(fecha)


def _area_elegida(request):
    area = request.GET.get('area', '')
    return Area.objects.filter(pk=area).first() if area.isdigit() else None


def _celdas(horas):
    """(valor, clase css) de cada columna de horas, en orden."""
    return [(horas[c], CLASES.get(c, '')) for c in CLAVES]


# Columnas resaltadas en la tabla y en el Excel.
CLASES = {'recargo_nocturno': 'recargo', 'recargo_nocturno_festivo': 'recargo', 'total_horas': 'total'}


@administrador_requerido
def conteo(request):
    """Conteo de horas de una semana (lunes a domingo) con las columnas del formato de nómina."""
    lunes, domingo = _semana_elegida(request)
    area = _area_elegida(request)
    filas, totales, total_general = conteo_semana(lunes, area)

    # Filas agrupadas por funcionario, con su total al final. Los valores van en el orden de las columnas.
    grupos = []
    for fila in filas:
        if not grupos or grupos[-1]['funcionario'] != fila['funcionario']:
            total = totales[fila['funcionario']]
            grupos.append({'funcionario': fila['funcionario'], 'filas': [], 'total': _celdas(total)})
        grupos[-1]['filas'].append({**fila, 'valores': _celdas(fila['horas'])})

    return render(request, 'reportes/conteo.html', {
        'lunes': lunes,
        'domingo': domingo,
        'anterior': lunes - datetime.timedelta(days=7),
        'siguiente': lunes + datetime.timedelta(days=7),
        'hoy': timezone.localdate(),
        'areas': Area.objects.filter(activo=True),
        'area': area,
        'grupos': grupos,
        'columnas': COLUMNAS_HORAS,
        'total_general': _celdas(total_general),
        'inicio_noche': INICIO_NOCHE,
        'fin_noche': FIN_NOCHE,
    })


@administrador_requerido
def conteo_excel(request):
    """El mismo conteo en Excel, con el formato de la institución y totales con fórmulas."""
    lunes, domingo = _semana_elegida(request)
    area = _area_elegida(request)
    filas, _, _ = conteo_semana(lunes, area)

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Conteo de horas'
    fuente = lambda **kw: Font(name='Arial', **kw)
    borde = Border(*(Side(style='thin', color='999999'),) * 4)

    hoja['A1'] = f'Conteo de horas · semana del {lunes:%d/%m/%Y} al {domingo:%d/%m/%Y}' + (f' · {area}' if area else '')
    hoja['A1'].font = fuente(bold=True, size=13)
    hoja['A2'] = (f'Nocturna: {INICIO_NOCHE:%H:%M} a {FIN_NOCHE:%H:%M}. Festivo: domingos y festivos. '
                  'Horas extra: solicitudes aprobadas. Recargos, diurna festivo y nocturnas festivas: turnos asignados. '
                  'Total horas: turnos más horas extra aprobadas.')
    hoja['A2'].font = fuente(italic=True, size=9, color='666666')

    encabezados = [(None, t) for t in ('FECHA', 'NOMBRE', 'SECCION', 'ACTIVIDAD')] + COLUMNAS_HORAS
    fila_encabezado = 4
    for col, (clave, texto) in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila_encabezado, column=col, value=texto)
        celda.font = fuente(bold=True)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        celda.border = borde
        celda.fill = PatternFill('solid', fgColor='D9D9D9' if CLASES.get(clave) else 'FFFFFF')
    hoja.row_dimensions[fila_encabezado].height = 45
    for col, ancho in enumerate([12, 32, 22, 30] + [14] * len(COLUMNAS_HORAS), start=1):
        hoja.column_dimensions[get_column_letter(col)].width = ancho

    primera = fila_encabezado + 1
    for i, fila in enumerate(filas):
        r = primera + i
        valores = [fila['fecha'], fila['funcionario'].nombre_completo, fila['seccion'], fila['actividad']]
        valores += [fila['horas'][clave] for clave in CLAVES]
        for col, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=r, column=col, value=valor)
            celda.font = fuente()
            celda.border = borde
            if col == 1:
                celda.number_format = 'DD/MM/YYYY'
            elif col > 4:
                celda.number_format = '0.00;-0.00;-'
    ultima = primera + len(filas) - 1

    fila_total = ultima + 1 if filas else primera
    hoja.cell(row=fila_total, column=4, value='TOTAL').font = fuente(bold=True)
    for col in range(5, 5 + len(COLUMNAS_HORAS)):
        letra = get_column_letter(col)
        celda = hoja.cell(row=fila_total, column=col, value=f'=SUM({letra}{primera}:{letra}{max(ultima, primera)})' if filas else 0)
        celda.font = fuente(bold=True)
        celda.number_format = '0.00;-0.00;-'
        celda.border = borde
    hoja.freeze_panes = hoja.cell(row=primera, column=1)

    salida = io.BytesIO()
    libro.save(salida)
    respuesta = HttpResponse(
        salida.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    respuesta['Content-Disposition'] = f'attachment; filename="conteo_horas_{lunes:%Y-%m-%d}.xlsx"'
    return respuesta
