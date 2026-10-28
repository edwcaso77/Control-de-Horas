"""Lectura de cuadros de turno en Excel con el formato de la institución (como «Enfermero Profesional.xlsx").

Cada hoja es un cuadro: un periodo («OCTUBRE 2026»), una fila de encabezado con «FECHAS / DIAS» y los días 1, 2, 3…,
y debajo un funcionario por fila (o dos filas con el nombre combinado, para dos turnos el mismo día).

Las celdas se leen así:
- La abreviatura de un tipo de turno (C1, M2, HC…). Con una F delante (FC1) es el mismo turno en festivo.
- Un número N: extramural (E) de N horas desde las 5:00 p. m.
- V: vacaciones, no se carga.
- Los turnos de horas variables sin número (HC) llevan sus horas mínimas.
"""
import datetime
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal

from openpyxl import load_workbook

from funcionarios.models import Funcionario

from .models import TURNOS_POR_DIA, TipoTurno

MESES = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE',
         'NOVIEMBRE', 'DICIEMBRE']
PERIODO = re.compile(r'\b(' + '|'.join(MESES) + r'|SETIEMBRE)\s+(?:DE\s+|DEL\s+)?(\d{4})\b')
FILAS_ENCABEZADO = 20  # el encabezado «FECHAS / DIAS» se busca en las primeras filas
EXTRAMURAL = 'E'
INICIO_EXTRAMURAL = datetime.time(17, 0)
VACACIONES = {'V', 'VAC', 'VACACIONES'}


class ErrorExcel(Exception):
    pass


def normalizar(texto):
    """Mayúsculas, sin tildes y con un solo espacio entre palabras, para comparar nombres."""
    texto = unicodedata.normalize('NFKD', str(texto))
    texto = ''.join(c for c in texto if not unicodedata.combining(c))
    return ' '.join(texto.upper().split())


def _entero(valor):
    if isinstance(valor, bool):
        return None
    if isinstance(valor, int):
        return valor
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return None


def _buscar_periodo(hoja, hasta_fila):
    for fila in hoja.iter_rows(min_row=1, max_row=hasta_fila, max_col=min(hoja.max_column, 60), values_only=True):
        for valor in fila:
            if isinstance(valor, (datetime.date, datetime.datetime)):
                return valor.year, valor.month
            if isinstance(valor, str):
                coincide = PERIODO.search(normalizar(valor))
                if coincide:
                    mes = 9 if coincide.group(1) == 'SETIEMBRE' else MESES.index(coincide.group(1)) + 1
                    return int(coincide.group(2)), mes
    return None, None


def _buscar_encabezado(hoja):
    """(fila, {columna: día}) de la fila «FECHAS / DIAS». Los días son la serie 1, 2, 3… que sigue a la derecha."""
    for fila in range(1, min(hoja.max_row, FILAS_ENCABEZADO) + 1):
        for columna in range(1, min(hoja.max_column, 60) + 1):
            valor = hoja.cell(row=fila, column=columna).value
            if not (isinstance(valor, str) and 'FECHA' in normalizar(valor)):
                continue
            dias, esperado = {}, 1
            for c in range(columna + 1, hoja.max_column + 1):
                numero = _entero(hoja.cell(row=fila, column=c).value)
                if numero == esperado:
                    dias[c] = numero
                    esperado += 1
                elif dias:
                    break
            if len(dias) >= 28:
                return fila, dias
    return None, {}


def _filas_combinadas(hoja):
    """{(fila, columna): fila de arriba} de las celdas combinadas, para saber qué filas son del mismo funcionario."""
    arriba = {}
    for rango in hoja.merged_cells.ranges:
        for fila in range(rango.min_row + 1, rango.max_row + 1):
            for columna in range(rango.min_col, rango.max_col + 1):
                arriba[(fila, columna)] = rango.min_row
    return arriba


class Catalogo:
    """Tipos de turno activos y funcionarios, para traducir códigos y nombres del Excel."""

    def __init__(self):
        self.tipos = {t.abreviatura.upper(): t for t in TipoTurno.objects.filter(activo=True)}
        self.personas = defaultdict(list)
        for f in Funcionario.objects.select_related('usuario'):
            for nombre in {normalizar(f'{f.nombres} {f.apellidos}'), normalizar(f'{f.apellidos} {f.nombres}')}:
                if f not in self.personas[nombre]:
                    self.personas[nombre].append(f)

    def funcionario(self, nombre):
        """El funcionario con ese nombre, o None si no hay uno solo."""
        encontrados = self.personas.get(normalizar(nombre), [])
        activos = [f for f in encontrados if f.activo]
        if len(activos) == 1:
            return activos[0]
        return encontrados[0] if len(encontrados) == 1 else None

    def turno(self, valor):
        """(tipo, horas, hora de inicio) de una celda; 'vacaciones'; o None si el código no se conoce."""
        numero = valor if isinstance(valor, (int, float)) and not isinstance(valor, bool) else None
        texto = normalizar(valor).replace(' ', '')
        if numero is None and re.fullmatch(r'\d+([.,]\d+)?', texto):
            numero = float(texto.replace(',', '.'))
        if numero is not None:
            extramural = self.tipos.get(EXTRAMURAL)
            if extramural is None or numero <= 0:
                return None
            return extramural, Decimal(str(numero)).normalize(), INICIO_EXTRAMURAL
        if texto in VACACIONES:
            return 'vacaciones'
        tipo = self.tipos.get(texto)
        if tipo is None and texto.startswith('F'):
            tipo = self.tipos.get(texto[1:])  # F = el turno cae en festivo
        if tipo is None:
            return None
        if tipo.horas_variables:
            return tipo, tipo.horas_minimas, INICIO_EXTRAMURAL if tipo.pide_hora_inicio else None
        return tipo, None, None


def leer_hoja(hoja, catalogo):
    """Resumen de una hoja y sus turnos: {funcionario_id: {día: [(tipo_id, horas, inicio), …]}}."""
    resultado = {
        'hoja': hoja.title, 'anio': None, 'mes': None, 'turnos': {}, 'personas': [], 'no_encontrados': [],
        'codigos_desconocidos': {}, 'vacaciones': 0, 'avisos': [], 'total': 0,
    }
    fila_encabezado, dias = _buscar_encabezado(hoja)
    if not dias:
        resultado['avisos'].append('No se encontró la fila «FECHAS / DIAS» con los días del mes.')
        return resultado
    resultado['anio'], resultado['mes'] = _buscar_periodo(hoja, fila_encabezado)
    columna_nombre = min(dias) - 1
    combinadas = _filas_combinadas(hoja)

    desconocidos = Counter()
    por_persona = {}  # fila donde empieza cada funcionario -> datos
    actual = None
    for fila in range(fila_encabezado + 1, hoja.max_row + 1):
        nombre = hoja.cell(row=fila, column=columna_nombre).value
        # Los nombres no llevan números: así se saltan filas de convenciones como «5:00 AM 7:00 AM».
        if isinstance(nombre, str) and nombre.strip() and not re.search(r'\d', nombre):
            actual = por_persona.setdefault(fila, {'nombre': ' '.join(nombre.split()), 'dias': defaultdict(list)})
        elif combinadas.get((fila, columna_nombre)) not in por_persona:
            actual = None
        if actual is None:
            continue
        for columna, dia in dias.items():
            valor = hoja.cell(row=fila, column=columna).value
            if valor is None or (isinstance(valor, str) and not valor.strip()):
                continue
            turno = catalogo.turno(valor)
            if turno == 'vacaciones':
                resultado['vacaciones'] += 1
            elif turno is None:
                desconocidos[normalizar(valor)] += 1
            else:
                actual['dias'][dia].append(turno)

    for datos in por_persona.values():
        if not datos['dias']:
            continue  # filas sin turnos reconocidos (títulos, convenciones…)
        funcionario = catalogo.funcionario(datos['nombre'])
        if funcionario is None:
            resultado['no_encontrados'].append(datos['nombre'])
            continue
        turnos = resultado['turnos'].setdefault(str(funcionario.pk), {})
        for dia, lista in datos['dias'].items():
            unicos = list({tipo.pk: (tipo, horas, inicio) for tipo, horas, inicio in lista}.values())
            if len(unicos) > TURNOS_POR_DIA:
                resultado['avisos'].append(
                    f'{funcionario.nombre_completo}, día {dia}: tiene {len(unicos)} turnos; se cargan solo los '
                    f'{TURNOS_POR_DIA} primeros ({", ".join(t.abreviatura for t, _, _ in unicos[:TURNOS_POR_DIA])}).'
                )
                unicos = unicos[:TURNOS_POR_DIA]
            turnos.setdefault(str(dia), []).extend(
                [tipo.pk, str(horas) if horas is not None else None, inicio.strftime('%H:%M') if inicio else None]
                for tipo, horas, inicio in unicos
            )
        cantidad = sum(len(v) for v in turnos.values())
        resultado['personas'].append({'nombre': funcionario.nombre_completo, 'turnos': cantidad})
        resultado['total'] += cantidad
    resultado['codigos_desconocidos'] = dict(desconocidos.most_common())
    return resultado


def leer_excel(archivo):
    """Lee todas las hojas del libro y devuelve el resumen de cada una que parezca un cuadro de turno."""
    try:
        libro = load_workbook(archivo, data_only=True)
    except Exception as error:  # archivo dañado o que no es .xlsx
        raise ErrorExcel('No se pudo abrir el archivo. Debe ser un Excel .xlsx.') from error
    catalogo = Catalogo()
    hojas = [leer_hoja(hoja, catalogo) for hoja in libro.worksheets]
    return [h for h in hojas if h['total'] or h['no_encontrados'] or h['codigos_desconocidos']]


def fechas_de(hoja, anio, mes):
    """Turnos de una hoja leída, ya con fecha: {(funcionario_id, fecha): [(tipo_id, horas, inicio)]}. Ignora días que el mes no tiene."""
    resultado = {}
    for funcionario_id, dias in hoja['turnos'].items():
        for dia, turnos in dias.items():
            try:
                fecha = datetime.date(anio, mes, int(dia))
            except ValueError:
                continue
            resultado[(int(funcionario_id), fecha)] = [
                (tipo_id, Decimal(horas) if horas else None, datetime.time.fromisoformat(inicio) if inicio else None)
                for tipo_id, horas, inicio in turnos
            ]
    return resultado

