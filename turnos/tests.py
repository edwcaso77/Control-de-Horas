import datetime
import io
from decimal import Decimal

from django.test import TestCase, override_settings
from django.urls import reverse

from funcionarios.models import Area, Cargo, Funcionario
from funcionarios.tests import crear_funcionario

from .forms import TipoTurnoForm
from .models import AsignacionTurno, TipoTurno


class TipoTurnoTests(TestCase):
    def test_duracion_y_horario(self):
        noche = TipoTurno(nombre='Noche', abreviatura='N', hora_inicio=datetime.time(19), hora_fin=datetime.time(7))
        self.assertEqual(noche.duracion_horas, 12)
        self.assertEqual(noche.horario, '19:00 – 07:00')
        descanso = TipoTurno(nombre='Descanso', abreviatura='D')
        self.assertEqual(descanso.duracion_horas, 0)
        self.assertEqual(descanso.horario, 'Sin horario')

    def test_color_de_texto_legible(self):
        self.assertEqual(TipoTurno(color='#1F2A5C').color_texto, '#FFFFFF')
        self.assertEqual(TipoTurno(color='#6EC6D9').color_texto, '#333333')


class DatosCuadroMixin:
    """Administrador, Ana (Urgencias), Luis (Hospitalización) y los turnos Mañana y Noche."""

    def setUp(self):
        self.admin = crear_funcionario('1000', es_administrador=True)
        self.client.force_login(self.admin.usuario)
        self.urgencias = Area.objects.get(nombre='Urgencias')
        self.otra_area = Area.objects.create(nombre='Hospitalización')
        self.ana = crear_funcionario('2000', nombres='Ana')
        self.luis = crear_funcionario('3000', nombres='Luis', area=self.otra_area.pk)
        self.manana = TipoTurno.objects.create(
            nombre='Mañana', abreviatura='M', hora_inicio=datetime.time(7), hora_fin=datetime.time(13), color='#6EC6D9',
        )
        self.noche = TipoTurno.objects.create(
            nombre='Noche', abreviatura='N', hora_inicio=datetime.time(19), hora_fin=datetime.time(7), color='#1F2A5C',
        )
        self.fecha = datetime.date(2026, 10, 21)  # miércoles


class CuadroTurnoTests(DatosCuadroMixin, TestCase):
    def _url_dia(self, fecha=None):
        fecha = fecha or self.fecha
        return reverse('turnos:dia', args=[fecha.year, fecha.month, fecha.day])

    def test_crear_tipo_de_turno(self):
        respuesta = self.client.post(reverse('turnos:tipo_crear'), {
            'nombre': 'Tarde', 'abreviatura': 't', 'hora_inicio': '13:00', 'hora_fin': '19:00', 'color': '#66b32e',
        })
        self.assertRedirects(respuesta, reverse('turnos:tipos'))
        tarde = TipoTurno.objects.get(nombre='Tarde')
        self.assertEqual((tarde.abreviatura, tarde.color), ('T', '#66B32E'))

    def test_tipo_con_una_sola_hora_es_invalido(self):
        respuesta = self.client.post(reverse('turnos:tipo_crear'), {
            'nombre': 'Raro', 'abreviatura': 'R', 'hora_inicio': '13:00', 'hora_fin': '', 'color': '#000000',
        })
        self.assertEqual(respuesta.status_code, 200)
        self.assertFalse(TipoTurno.objects.filter(nombre='Raro').exists())

    def test_asignar_turnos_de_un_dia(self):
        respuesta = self.client.post(self._url_dia(), {
            f'turno_{self.ana.pk}_1': self.manana.pk,
            f'turno_{self.luis.pk}_1': self.noche.pk,
            'alcance': 'dia',
        })
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=10")
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.ana, fecha=self.fecha).tipo_turno, self.manana)
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.luis, fecha=self.fecha).tipo_turno, self.noche)
        self.assertEqual(AsignacionTurno.objects.count(), 2)

    def test_aplicar_a_toda_la_semana(self):
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': self.noche.pk, 'alcance': 'semana'})
        fechas = list(AsignacionTurno.objects.filter(funcionario=self.ana).values_list('fecha', flat=True))
        self.assertEqual(fechas, [datetime.date(2026, 10, d) for d in range(19, 26)])

    def test_cambiar_y_quitar_turno(self):
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': self.manana.pk, 'alcance': 'dia'})
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': self.noche.pk, 'alcance': 'dia'})
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.ana).tipo_turno, self.noche)
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': '', 'alcance': 'dia'})
        self.assertFalse(AsignacionTurno.objects.filter(funcionario=self.ana).exists())

    def test_filtro_por_area_solo_toca_esa_area(self):
        AsignacionTurno.objects.create(funcionario=self.luis, fecha=self.fecha, tipo_turno=self.noche)
        respuesta = self.client.get(self._url_dia(), {'area': self.urgencias.pk})
        nombres = [fila['funcionario'].nombres for fila in respuesta.context['filas']]
        self.assertNotIn('Luis', nombres)

        # Guardar Urgencias sin turno no borra el turno de Luis, que es de otra área.
        self.client.post(self._url_dia(), {'area': self.urgencias.pk, f'turno_{self.ana.pk}_1': '', 'alcance': 'dia'})
        self.assertTrue(AsignacionTurno.objects.filter(funcionario=self.luis).exists())

    def test_filtro_por_cargo(self):
        medico = Cargo.objects.create(nombre='Médico')
        sofia = crear_funcionario('4000', nombres='Sofía', cargo=medico.pk)
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.manana)
        AsignacionTurno.objects.create(funcionario=sofia, fecha=self.fecha, tipo_turno=self.noche)

        # El cuadro del mes solo muestra los turnos de ese cargo, y lo conserva al cambiar de mes y al abrir un día.
        respuesta = self.client.get(reverse('turnos:cuadro'), {'anio': 2026, 'mes': 10, 'cargo': medico.pk})
        dia = next(d for s in respuesta.context['semanas'] for d in s if d['fecha'] == self.fecha)
        self.assertEqual([a.funcionario for a in dia['turnos']], [sofia])
        self.assertContains(respuesta, f'?cargo={medico.pk}')

        respuesta = self.client.get(self._url_dia(), {'cargo': medico.pk})
        self.assertEqual([fila['funcionario'] for fila in respuesta.context['filas']], [sofia])

        # Guardar el día con el filtro no toca a los de otro cargo y vuelve al cuadro filtrado.
        respuesta = self.client.post(self._url_dia(), {'cargo': medico.pk, f'turno_{sofia.pk}_1': '', 'alcance': 'dia'})
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=10&cargo={medico.pk}")
        self.assertTrue(AsignacionTurno.objects.filter(funcionario=self.ana).exists())
        self.assertFalse(AsignacionTurno.objects.filter(funcionario=sofia).exists())

    def test_turno_inexistente_se_ignora(self):
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': '99999', 'alcance': 'dia'})
        self.assertFalse(AsignacionTurno.objects.exists())

    def test_cada_dia_lista_nombres_ordenados_por_turno(self):
        AsignacionTurno.objects.create(funcionario=self.luis, fecha=self.fecha, tipo_turno=self.noche)
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.manana)
        respuesta = self.client.get(reverse('turnos:cuadro'), {'anio': 2026, 'mes': 10})
        dia = next(d for s in respuesta.context['semanas'] for d in s if d['fecha'] == self.fecha)
        # Mañana (07:00) va antes que Noche (19:00) porque los tipos se ordenan por hora de inicio.
        self.assertEqual([(a.tipo_turno.abreviatura, a.funcionario.nombre_corto) for a in dia['turnos']],
                         [('M', 'Ana Pérez'), ('N', 'Luis Pérez')])
        self.assertContains(respuesta, 'Luis Pérez')
        self.assertNotContains(respuesta, 'Mañana ·')

    def test_funcionario_inactivo_no_aparece(self):
        self.ana.usuario.is_active = False
        self.ana.usuario.save()
        filas = self.client.get(self._url_dia()).context['filas']
        self.assertNotIn(self.ana, [f['funcionario'] for f in filas])

    def test_fecha_invalida(self):
        self.assertEqual(self.client.get('/cuadro-de-turno/dia/2026/2/30/').status_code, 404)

    def test_solo_administrador(self):
        self.client.force_login(self.ana.usuario)
        self.assertEqual(self.client.get(reverse('turnos:cuadro')).status_code, 403)
        self.assertEqual(self.client.get(self._url_dia()).status_code, 403)
        self.assertEqual(self.client.get(reverse('turnos:tipos')).status_code, 403)

    def test_menu_lateral_muestra_cuadro_de_turno_solo_al_admin(self):
        self.assertContains(self.client.get(reverse('calendario:mes')), 'Cuadro de turno')
        self.client.force_login(self.ana.usuario)
        self.assertNotContains(self.client.get(reverse('calendario:mes')), 'Cuadro de turno')


class VariosDiasTests(DatosCuadroMixin, TestCase):
    fechas = [datetime.date(2026, 10, 20), datetime.date(2026, 10, 22), datetime.date(2026, 10, 27)]

    def _url(self):
        return reverse('turnos:varios')

    def _texto(self, fechas=None):
        return ','.join(d.isoformat() for d in (fechas or self.fechas))

    def _celda(self, funcionario, fecha):
        return f'turno_{funcionario.pk}_{fecha.isoformat()}_1'

    def test_cada_celda_se_asigna_por_separado(self):
        d1, d2, d3 = self.fechas
        respuesta = self.client.post(self._url(), {
            'fechas': self._texto(),
            self._celda(self.ana, d1): self.manana.pk,
            self._celda(self.ana, d2): self.noche.pk,
            self._celda(self.ana, d3): '',
            self._celda(self.luis, d2): self.manana.pk,
        })
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=10")
        turnos = {(a.funcionario_id, a.fecha): a.tipo_turno for a in AsignacionTurno.objects.all()}
        self.assertEqual(turnos, {
            (self.ana.pk, d1): self.manana,
            (self.ana.pk, d2): self.noche,
            (self.luis.pk, d2): self.manana,
        })

    def test_celda_vacia_quita_y_celda_ausente_conserva(self):
        d1, d2, _ = self.fechas
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=d1, tipo_turno=self.noche)
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=d2, tipo_turno=self.noche)
        self.client.post(self._url(), {'fechas': self._texto(), self._celda(self.ana, d1): ''})
        self.assertEqual(list(AsignacionTurno.objects.values_list('fecha', flat=True)), [d2])

    def test_celda_sin_cambios_no_modifica_quien_asigno(self):
        d1 = self.fechas[0]
        asignacion = AsignacionTurno.objects.create(funcionario=self.ana, fecha=d1, tipo_turno=self.noche)
        self.client.post(self._url(), {'fechas': self._texto(), self._celda(self.ana, d1): self.noche.pk})
        asignacion.refresh_from_db()
        self.assertIsNone(asignacion.asignado_por)

    def test_tabla_muestra_turnos_actuales_por_dia(self):
        AsignacionTurno.objects.create(funcionario=self.luis, fecha=self.fechas[1], tipo_turno=self.manana)
        filas = {f['funcionario'].pk: f for f in self.client.get(self._url(), {'fechas': self._texto()}).context['filas']}
        principales = [c['ranuras'][0] for c in filas[self.luis.pk]['celdas']]
        self.assertEqual([a.tipo_turno if a else None for a in principales], [None, self.manana, None])

    def test_filtro_por_cargo(self):
        enfermero = Cargo.objects.create(nombre='Enfermero')
        self.luis.cargo = enfermero
        self.luis.save()
        respuesta = self.client.get(self._url(), {'fechas': self._texto(), 'cargo': enfermero.pk})
        self.assertEqual([f['funcionario'] for f in respuesta.context['filas']], [self.luis])

        # Guardar con el filtro de cargo no toca a funcionarios de otros cargos.
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fechas[0], tipo_turno=self.noche)
        self.client.post(self._url(), {
            'fechas': self._texto(), 'cargo': enfermero.pk, self._celda(self.ana, self.fechas[0]): '',
        })
        self.assertTrue(AsignacionTurno.objects.filter(funcionario=self.ana).exists())

    def test_fechas_invalidas_o_vacias(self):
        respuesta = self.client.get(self._url(), {'fechas': 'hola,2026-13-40'})
        self.assertRedirects(respuesta, reverse('turnos:cuadro'))
        contexto = self.client.get(self._url(), {'fechas': '2026-10-22,basura,2026-10-20,2026-10-22'}).context
        self.assertEqual(contexto['fechas'], [datetime.date(2026, 10, 20), datetime.date(2026, 10, 22)])

    def test_respeta_el_area(self):
        AsignacionTurno.objects.create(funcionario=self.luis, fecha=self.fechas[0], tipo_turno=self.noche)
        self.client.post(self._url(), {
            'fechas': self._texto(), 'area': self.urgencias.pk, self._celda(self.luis, self.fechas[0]): '',
        })
        self.assertTrue(AsignacionTurno.objects.filter(funcionario=self.luis).exists())

    def test_solo_administrador_en_varios_dias(self):
        self.client.force_login(self.ana.usuario)
        self.assertEqual(self.client.get(self._url(), {'fechas': self._texto()}).status_code, 403)

    def test_cuadro_tiene_casillas_y_barra(self):
        respuesta = self.client.get(reverse('turnos:cuadro'), {'anio': 2026, 'mes': 10})
        self.assertContains(respuesta, 'class="casilla-dia" value="2026-10-20"')
        self.assertContains(respuesta, 'id="barra-seleccion"')


class BorrarAgendaTests(DatosCuadroMixin, TestCase):
    def setUp(self):
        super().setUp()
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.manana)
        AsignacionTurno.objects.create(funcionario=self.luis, fecha=self.fecha, tipo_turno=self.noche)

    def test_en_produccion_no_aparece_ni_funciona(self):
        # Las pruebas corren con DEBUG=False, igual que el sitio publicado.
        self.assertNotContains(self.client.get(reverse('turnos:cuadro')), 'Eliminar toda la agenda')
        self.assertEqual(self.client.post(reverse('turnos:borrar_agenda')).status_code, 404)
        self.assertEqual(AsignacionTurno.objects.count(), 2)

    @override_settings(DEBUG=True)
    def test_en_local_borra_solo_los_turnos(self):
        self.assertContains(self.client.get(reverse('turnos:cuadro')), 'Eliminar toda la agenda')
        respuesta = self.client.post(reverse('turnos:borrar_agenda'), {'anio': 2026, 'mes': 10})
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=10")
        self.assertEqual(AsignacionTurno.objects.count(), 0)
        self.assertEqual(TipoTurno.objects.count(), 2)
        self.assertTrue(Funcionario.objects.filter(pk=self.ana.pk).exists())

    @override_settings(DEBUG=True)
    def test_solo_administrador_y_por_post(self):
        self.assertEqual(self.client.get(reverse('turnos:borrar_agenda')).status_code, 405)
        self.client.force_login(self.ana.usuario)
        self.assertEqual(self.client.post(reverse('turnos:borrar_agenda')).status_code, 403)
        self.assertEqual(AsignacionTurno.objects.count(), 2)


class TurnosEspecialesTests(DatosCuadroMixin, TestCase):
    """Jornada partida (A), turnos de horas variables (HC, extramural) y dos turnos el mismo día."""

    def setUp(self):
        super().setUp()
        self.partido = TipoTurno.objects.create(
            nombre='Jornada partida', abreviatura='A',
            hora_inicio=datetime.time(8), hora_fin=datetime.time(12),
            hora_inicio_2=datetime.time(14), hora_fin_2=datetime.time(18),
        )
        self.hc = TipoTurno.objects.create(
            nombre='Disponibilidad', abreviatura='HC', horas_variables=True,
            horas_minimas=Decimal('2'), horas_maximas=Decimal('4'),
        )
        self.extramural = TipoTurno.objects.create(
            nombre='Extramural', abreviatura='1-12', horas_variables=True,
            horas_minimas=Decimal('1'), horas_maximas=Decimal('12'),
        )

    def _url_dia(self):
        return reverse('turnos:dia', args=[self.fecha.year, self.fecha.month, self.fecha.day])

    def test_jornada_partida_suma_los_dos_bloques(self):
        self.assertEqual(self.partido.duracion_horas, 8)
        self.assertEqual(self.partido.horario, '08:00 – 12:00 / 14:00 – 18:00')

    def test_turno_de_cuatro_bloques(self):
        base = {'nombre': 'A1 (especial)', 'abreviatura': 'A1', 'color': '#000000',
                'hora_inicio': '05:00', 'hora_fin': '07:00', 'hora_inicio_2': '10:30', 'hora_fin_2': '13:00',
                'hora_inicio_3': '16:00', 'hora_fin_3': '18:00'}
        self.assertFalse(TipoTurnoForm({**base, 'hora_inicio_3': '', 'hora_fin_3': '',
                                        'hora_inicio_4': '22:00', 'hora_fin_4': '00:00'}).is_valid())
        self.assertFalse(TipoTurnoForm({**base, 'hora_fin_3': ''}).is_valid())
        form = TipoTurnoForm({**base, 'hora_inicio_4': '22:00', 'hora_fin_4': '00:00'})
        self.assertTrue(form.is_valid(), form.errors)
        tipo = form.save()
        self.assertEqual(tipo.duracion_horas, 8.5)
        self.assertEqual(tipo.horario, '05:00 – 07:00 / 10:30 – 13:00 / 16:00 – 18:00 / 22:00 – 00:00')
        self.assertEqual(tipo.hora_fin_ultimo_bloque, datetime.time(0))

    def test_turno_de_horas_variables(self):
        self.assertEqual(self.hc.horario, 'Horas variables (2–4 h)')
        self.assertEqual(self.hc.duracion_horas, 0)
        asignacion = AsignacionTurno(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.hc, horas=Decimal('3'))
        self.assertEqual((asignacion.etiqueta, asignacion.horario, asignacion.horas_trabajadas), ('HC 3h', '3 h', 3.0))

    def test_formulario_valida_las_clases_de_turno(self):
        base = {'nombre': 'Nuevo', 'abreviatura': 'NV', 'color': '#000000'}
        self.assertFalse(TipoTurnoForm({**base, 'horas_variables': 'on', 'horas_minimas': '5', 'horas_maximas': '2'}).is_valid())
        self.assertFalse(TipoTurnoForm({**base, 'horas_variables': 'on'}).is_valid())
        self.assertFalse(TipoTurnoForm({**base, 'horas_variables': 'on', 'horas_minimas': '1', 'horas_maximas': '2',
                                        'hora_inicio': '08:00', 'hora_fin': '10:00'}).is_valid())
        self.assertFalse(TipoTurnoForm({**base, 'hora_inicio_2': '14:00', 'hora_fin_2': '18:00'}).is_valid())
        form = TipoTurnoForm({**base, 'hora_inicio': '08:00', 'hora_fin': '12:00', 'hora_inicio_2': '14:00', 'hora_fin_2': '18:00',
                              'horas_minimas': '1', 'horas_maximas': '3'})
        self.assertTrue(form.is_valid(), form.errors)
        tipo = form.save()
        # El rango de horas solo aplica a turnos de horas variables, así que no se guarda.
        self.assertIsNone(tipo.horas_minimas)

    def test_dos_turnos_el_mismo_dia_con_horas(self):
        respuesta = self.client.post(self._url_dia(), {
            f'turno_{self.ana.pk}_1': self.manana.pk,
            f'turno_{self.ana.pk}_2': self.hc.pk,
            f'horas_turno_{self.ana.pk}_2': '3',
            'alcance': 'dia',
        })
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=10")
        turnos = {a.tipo_turno.abreviatura: a.horas for a in AsignacionTurno.objects.filter(funcionario=self.ana)}
        self.assertEqual(turnos, {'M': None, 'HC': Decimal('3.0')})

        # Quitar el adicional deja solo el principal.
        self.client.post(self._url_dia(), {f'turno_{self.ana.pk}_1': self.manana.pk, f'turno_{self.ana.pk}_2': '', 'alcance': 'dia'})
        self.assertEqual(list(AsignacionTurno.objects.filter(funcionario=self.ana).values_list('tipo_turno__abreviatura', flat=True)), ['M'])

    def test_horas_fuera_de_rango_no_cambian_nada(self):
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.noche)
        respuesta = self.client.post(self._url_dia(), {
            f'turno_{self.ana.pk}_1': self.hc.pk, f'horas_turno_{self.ana.pk}_1': '9', 'alcance': 'dia',
        }, follow=True)
        self.assertContains(respuesta, 'HC necesita entre 2 y 4 horas')
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.ana).tipo_turno, self.noche)

    def test_el_mismo_turno_en_ambas_ranuras_cuenta_una_vez(self):
        self.client.post(self._url_dia(), {
            f'turno_{self.ana.pk}_1': self.partido.pk, f'turno_{self.ana.pk}_2': self.partido.pk, 'alcance': 'dia',
        })
        self.assertEqual(AsignacionTurno.objects.filter(funcionario=self.ana).count(), 1)

    def test_cambiar_horas_actualiza_el_turno(self):
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.extramural, horas=Decimal('6'))
        self.client.post(self._url_dia(), {
            f'turno_{self.ana.pk}_1': self.extramural.pk, f'horas_turno_{self.ana.pk}_1': '10', 'alcance': 'dia',
        })
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.ana).horas, Decimal('10'))

    def test_varios_dias_con_turno_adicional(self):
        fecha = datetime.date(2026, 10, 22)
        clave = f'turno_{self.ana.pk}_{fecha.isoformat()}'
        self.client.post(reverse('turnos:varios'), {
            'fechas': fecha.isoformat(),
            f'{clave}_1': self.partido.pk, f'{clave}_2': self.extramural.pk, f'horas_{clave}_2': '5',
        })
        turnos = {a.tipo_turno.abreviatura: a.horas for a in AsignacionTurno.objects.filter(funcionario=self.ana, fecha=fecha)}
        self.assertEqual(turnos, {'A': None, '1-12': Decimal('5.0')})

    def test_listas_muestran_las_horas(self):
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.hc, horas=Decimal('3'))
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=self.fecha, tipo_turno=self.manana)
        cuadro = self.client.get(reverse('turnos:cuadro'), {'anio': 2026, 'mes': 10})
        self.assertContains(cuadro, '<b>HC 3h</b>')
        dia = next(d for s in cuadro.context['semanas'] for d in s if d['fecha'] == self.fecha)
        # Los turnos de horas variables (sin hora de inicio) quedan al final.
        self.assertEqual([a.tipo_turno.abreviatura for a in dia['turnos']], ['M', 'HC'])
        consulta = self.client.get(reverse('calendario:dia', args=[2026, 10, 21]))
        self.assertContains(consulta, '3 h')


class HoraDeInicioAlAsignarTests(DatosCuadroMixin, TestCase):
    """Extramural: además de las horas, se escribe a qué hora empieza cada vez."""

    def setUp(self):
        super().setUp()
        self.extramural = TipoTurno.objects.create(
            nombre='Extramural', abreviatura='1-12', horas_variables=True, pide_hora_inicio=True,
            horas_minimas=Decimal('1'), horas_maximas=Decimal('12'),
        )
        self.url = reverse('turnos:dia', args=[self.fecha.year, self.fecha.month, self.fecha.day])
        self.clave = f'turno_{self.ana.pk}'

    def test_guarda_hora_de_inicio_y_calcula_el_fin(self):
        self.client.post(self.url, {
            f'{self.clave}_1': self.extramural.pk, f'horas_{self.clave}_1': '6', f'inicio_{self.clave}_1': '07:00',
            'alcance': 'dia',
        })
        asignacion = AsignacionTurno.objects.get(funcionario=self.ana)
        self.assertEqual(asignacion.hora_inicio, datetime.time(7))
        self.assertEqual(asignacion.hora_fin, datetime.time(13))
        self.assertEqual(asignacion.horario, '07:00 – 13:00 (6 h)')

    def test_termina_al_dia_siguiente(self):
        asignacion = AsignacionTurno(
            funcionario=self.ana, fecha=self.fecha, tipo_turno=self.extramural, horas=Decimal('8'), hora_inicio=datetime.time(20),
        )
        self.assertEqual(asignacion.horario, '20:00 – 04:00 (8 h)')

    def test_sin_hora_de_inicio_no_se_guarda(self):
        respuesta = self.client.post(self.url, {
            f'{self.clave}_1': self.extramural.pk, f'horas_{self.clave}_1': '6', f'inicio_{self.clave}_1': '',
            'alcance': 'dia',
        }, follow=True)
        self.assertContains(respuesta, '1-12 necesita la hora de inicio')
        self.assertFalse(AsignacionTurno.objects.exists())

    def test_cambiar_solo_la_hora_de_inicio_actualiza(self):
        AsignacionTurno.objects.create(
            funcionario=self.ana, fecha=self.fecha, tipo_turno=self.extramural, horas=Decimal('6'), hora_inicio=datetime.time(7),
        )
        self.client.post(self.url, {
            f'{self.clave}_1': self.extramural.pk, f'horas_{self.clave}_1': '6', f'inicio_{self.clave}_1': '09:30',
            'alcance': 'dia',
        })
        self.assertEqual(AsignacionTurno.objects.get(funcionario=self.ana).hora_inicio, datetime.time(9, 30))

    def test_la_pantalla_muestra_la_hora_guardada(self):
        AsignacionTurno.objects.create(
            funcionario=self.ana, fecha=self.fecha, tipo_turno=self.extramural, horas=Decimal('6'), hora_inicio=datetime.time(7),
        )
        self.assertContains(self.client.get(self.url), f'name="inicio_{self.clave}_1" class="inicio-turno"\n           value="07:00"')
        consulta = self.client.get(reverse('calendario:dia', args=[2026, 10, 21]))
        self.assertContains(consulta, '07:00 – 13:00 (6 h)')

    def test_solo_turnos_de_horas_variables_piden_hora_de_inicio(self):
        tipo = TipoTurno(nombre='Fijo', abreviatura='F', hora_inicio=datetime.time(8), hora_fin=datetime.time(12), pide_hora_inicio=True)
        tipo.full_clean()
        self.assertFalse(tipo.pide_hora_inicio)


class SubirExcelTests(TestCase):
    """Carga del cuadro de turno desde un Excel con el formato de la institución."""

    def setUp(self):
        self.admin = crear_funcionario('1000', es_administrador=True, nombres='Admin', apellidos='Zuluaga')
        self.ana = crear_funcionario('2000', nombres='Ana María', apellidos='Pérez Gómez')
        self.m1 = TipoTurno.objects.create(nombre='Mañana', abreviatura='M1', hora_inicio=datetime.time(5), hora_fin=datetime.time(11))
        self.hc = TipoTurno.objects.create(nombre='Disponibilidad', abreviatura='HC', horas_variables=True,
                                           horas_minimas=Decimal('2'), horas_maximas=Decimal('4'))
        self.e = TipoTurno.objects.create(nombre='Extramural', abreviatura='E', horas_variables=True, pide_hora_inicio=True,
                                          horas_minimas=Decimal('1'), horas_maximas=Decimal('12'))
        self.client.force_login(self.admin.usuario)

    def _excel(self):
        from openpyxl import Workbook
        libro = Workbook()
        hoja = libro.active
        hoja.title = 'JEFES'
        hoja['B2'], hoja['C2'] = 'PERIODO:', 'OCTUBRE 2026'
        hoja['A4'] = 'FECHAS / DIAS'
        for dia in range(1, 32):
            hoja.cell(row=4, column=2 + dia, value=dia)
        hoja['B6'] = '  ANA MARIA PEREZ GOMEZ'  # sin tildes y con espacios de sobra
        hoja.merge_cells('B6:B7')
        hoja['C6'], hoja['C7'] = 'M1', 1         # día 1: M1 y extramural de 1 h
        hoja['D6'] = 'FM1'                       # día 2: M1 en festivo
        hoja['E6'] = 'HC'                        # día 3: disponibilidad (horas mínimas)
        hoja['F6'] = 'V'                         # día 4: vacaciones
        hoja['G6'] = 'ZZ'                        # día 5: código desconocido
        hoja['B8'] = 'PEDRO NO EXISTE'
        hoja['C8'] = 'M1'
        hoja['B10'] = '5:00 AM 7:00 AM'          # convención, no es un funcionario
        hoja['C10'] = 'M1'
        archivo = io.BytesIO()
        libro.save(archivo)
        archivo.seek(0)
        archivo.name = 'cuadro.xlsx'
        return archivo

    def test_vista_previa_y_carga(self):
        # Turno viejo en un día que el Excel deja vacío (vacaciones): se quita al cargar.
        AsignacionTurno.objects.create(funcionario=self.ana, fecha=datetime.date(2026, 10, 4), tipo_turno=self.m1)

        vista = self.client.post(reverse('turnos:subir_excel'), {'archivo': self._excel()})
        hoja = vista.context['hojas'][0]
        self.assertEqual((hoja['anio'], hoja['mes'], hoja['total']), (2026, 10, 4))
        self.assertEqual(hoja['no_encontrados'], ['PEDRO NO EXISTE'])
        self.assertEqual(hoja['codigos_desconocidos'], {'ZZ': 1})
        self.assertEqual(hoja['vacaciones'], 1)
        self.assertEqual(AsignacionTurno.objects.count(), 1)  # la vista previa no carga nada

        respuesta = self.client.post(reverse('turnos:cargar_excel'), {'hoja_0': 'on', 'mes_0': '10', 'anio_0': '2026', 'vaciar': 'on'})
        self.assertRedirects(respuesta, reverse('turnos:cuadro') + '?anio=2026&mes=10')
        turnos = {(a.fecha.day, a.tipo_turno.abreviatura): a for a in AsignacionTurno.objects.filter(funcionario=self.ana)}
        self.assertEqual(set(turnos), {(1, 'M1'), (1, 'E'), (2, 'M1'), (3, 'HC')})
        self.assertEqual((turnos[(1, 'E')].horas, turnos[(1, 'E')].hora_inicio), (Decimal('1'), datetime.time(17)))
        self.assertEqual(turnos[(3, 'HC')].horas, Decimal('2'))

        # Volver a cargar el mismo archivo no duplica nada.
        self.client.post(reverse('turnos:subir_excel'), {'archivo': self._excel()})
        self.client.post(reverse('turnos:cargar_excel'), {'hoja_0': 'on', 'mes_0': '10', 'anio_0': '2026', 'vaciar': 'on'})
        self.assertEqual(AsignacionTurno.objects.filter(funcionario=self.ana).count(), 4)

    def test_rechaza_archivos_que_no_son_excel(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        respuesta = self.client.post(reverse('turnos:subir_excel'), {'archivo': SimpleUploadedFile('cuadro.xlsx', b'no es excel')})
        self.assertContains(respuesta, 'No se pudo abrir el archivo')
        self.assertIsNone(respuesta.context['hojas'])

    def test_solo_administradores(self):
        self.client.force_login(self.ana.usuario)
        self.assertEqual(self.client.get(reverse('turnos:subir_excel')).status_code, 403)


class ProgramarMesTests(TestCase):
    """Cuatro titulares que en la última semana de octubre de 2026 tuvieron las líneas 0 a 3, un relevo y un comodín."""

    def setUp(self):
        from . import programar as rotacion
        self.rotacion = rotacion
        self.admin = crear_funcionario('1000', es_administrador=True)
        self.client.force_login(self.admin.usuario)
        self.cargo = Cargo.objects.create(nombre='Enfermero Profesional')
        horarios = {'C': (5, 14), 'X': (14, 23), 'M': (5, 11), 'T': (11, 17)}
        self.tipos = {}
        for codigo in rotacion.CODIGOS[:-1]:
            inicio, fin = horarios[codigo[0]]
            self.tipos[codigo] = TipoTurno.objects.create(
                nombre=codigo, abreviatura=codigo, hora_inicio=datetime.time(inicio), hora_fin=datetime.time(fin),
            )
        self.tipos['A'] = TipoTurno.objects.create(
            nombre='Administrativo', abreviatura='A', hora_inicio=datetime.time(8), hora_fin=datetime.time(12),
            hora_inicio_2=datetime.time(14), hora_fin_2=datetime.time(18),
        )
        nombres = ['Uno', 'Dos', 'Tres', 'Cuatro', 'Relevo', 'Comodin']
        self.personas = [
            crear_funcionario(str(2000 + i), nombres=nombre, apellidos=f'P{i}', cargo=self.cargo.pk)
            for i, nombre in enumerate(nombres)
        ]
        semana = [datetime.date(2026, 10, d) for d in range(26, 32)]  # lunes 26 a sábado 31
        for linea, persona in enumerate(self.personas[:4]):
            for fecha in semana:
                self._asignar(persona, fecha, rotacion.puesto(linea, fecha))
        for fecha, codigo in zip(semana, ['C1', 'T2', 'X1', 'M2']):
            self._asignar(self.personas[4], fecha, codigo)
        for fecha in semana[:5]:
            self._asignar(self.personas[5], fecha, 'A')

    def _asignar(self, persona, fecha, codigo):
        AsignacionTurno.objects.create(funcionario=persona, fecha=fecha, tipo_turno=self.tipos[codigo])

    def _grupo(self):
        return {p['funcionario'].pk: p for p in self.rotacion.grupo_del_mes(self.cargo, 2026, 11)}

    def test_grupo_sigue_la_rotacion_del_mes_anterior(self):
        grupo = self._grupo()
        for linea, persona in enumerate(self.personas[:4]):
            # Noviembre empieza en domingo: la primera semana ya es nueva y cada titular pasa a la línea siguiente.
            self.assertEqual((grupo[persona.pk]['papel'], grupo[persona.pk]['linea']), ('titular', (linea + 1) % 4))
        self.assertEqual(grupo[self.personas[4].pk]['papel'], 'relevo')
        self.assertEqual(grupo[self.personas[5].pk]['papel'], 'comodin')
        self.assertTrue(all(p['estado'] == 'sigue' for p in grupo.values()))

    def test_grupo_marca_quien_entra_y_quien_sale(self):
        uno = self.personas[0]
        uno.fecha_retiro = datetime.date(2026, 11, 15)
        uno.save()
        nueva = crear_funcionario('3000', nombres='Nueva', cargo=self.cargo.pk, fecha_ingreso='2026-11-10')
        grupo = self._grupo()
        self.assertEqual((grupo[uno.pk]['estado'], grupo[uno.pk]['hasta']), ('sale', datetime.date(2026, 11, 15)))
        self.assertEqual((grupo[nueva.pk]['estado'], grupo[nueva.pk]['desde']), ('entra', datetime.date(2026, 11, 10)))

    def test_inactivo_sin_retiro_ya_no_se_programa(self):
        uno = self.personas[0]
        uno.usuario.is_active = False
        uno.usuario.save()
        grupo = self._grupo()
        self.assertEqual((grupo[uno.pk]['estado'], grupo[uno.pk]['papel']), ('sale', 'fuera'))

    def test_borrador_cubre_los_cuatro_puestos_sin_pasar_la_jornada(self):
        grupo = self._grupo()
        personas = [self.rotacion.Persona(p['funcionario'], p['papel'], p['linea']) for p in grupo.values()]
        festivo = datetime.date(2026, 11, 2)
        borrador = self.rotacion.generar(personas, 2026, 11, self.tipos, {festivo})
        self.assertEqual((borrador.avisos, dict(borrador.vacantes)), ([], {}))
        for fecha in self.rotacion.dias_del_mes(2026, 11):
            del_dia = sorted(c for turnos in borrador.turnos.values() for f, c in turnos.items() if f == fecha and c != 'A')
            esperado = [] if fecha.weekday() == 6 else sorted(self.rotacion.puesto(l, fecha) for l in range(4))
            self.assertEqual(del_dia, esperado, fecha)
        for persona in self.personas[:5]:
            horas = self.rotacion.horas_por_semana(borrador, persona.pk, 2026, 11, self.tipos)
            self.assertLessEqual(max(horas), 42)
        # Cada titular descansa un día en las semanas completas, y el comodín no hace A el festivo.
        self.assertEqual(len(borrador.descansos[self.personas[0].pk]), 4)
        self.assertNotIn(festivo, borrador.turnos[self.personas[5].pk])

    def test_vacaciones_las_cubre_el_relevo(self):
        grupo = self._grupo()
        semana = {datetime.date(2026, 11, d) for d in range(9, 15)}
        personas = [
            self.rotacion.Persona(p['funcionario'], p['papel'], p['linea'], ausencias=semana if pk == self.personas[0].pk else set())
            for pk, p in grupo.items()
        ]
        borrador = self.rotacion.generar(personas, 2026, 11, self.tipos)
        self.assertFalse(semana & set(borrador.turnos[self.personas[0].pk]))
        self.assertEqual(dict(borrador.vacantes), {})

    def test_el_relevo_elige_dias_y_turnos(self):
        grupo = self._grupo()
        lunes, martes = datetime.date(2026, 11, 9), datetime.date(2026, 11, 10)
        semana = [lunes + datetime.timedelta(days=i) for i in range(6)]
        # Esa semana el relevo solo puede el lunes, y pide C1; el martes solo tarde (no le alcanza para descansos).
        disponibles = {d: set() for d in semana}
        disponibles[lunes] = self.rotacion.turnos_aceptados(lunes, 'C1')
        disponibles[martes] = self.rotacion.turnos_aceptados(martes, 'tarde')
        self.assertEqual(disponibles[martes], {'T1', 'T2'})
        personas = [
            self.rotacion.Persona(p['funcionario'], p['papel'], p['linea'],
                                  disponibles=disponibles if p['papel'] == 'relevo' else None)
            for p in grupo.values()
        ]
        borrador = self.rotacion.generar(personas, 2026, 11, self.tipos)
        relevo = borrador.turnos[self.personas[4].pk]
        self.assertEqual(relevo.get(lunes), 'C1')
        self.assertIn(relevo.get(martes), {'T1', 'T2', None})
        self.assertFalse({d for d in semana[2:]} & set(relevo))
        # Descansa el titular que tenía C1 ese lunes.
        quien = next(p for p in self.personas[:4] if lunes in borrador.descansos[p.pk])
        linea = (grupo[quien.pk]['linea'] + 1) % 4  # el 9 de noviembre es la segunda semana del mes
        self.assertEqual(self.rotacion.puesto(linea, lunes), 'C1')

    def test_guardar_el_borrador_reemplaza_los_turnos_del_mes(self):
        uno = self.personas[0]
        self._asignar(uno, datetime.date(2026, 11, 20), 'A')
        datos = {'cargo': self.cargo.pk, 'anio': 2026, 'mes': 11, 'accion': 'guardar'}
        for pk, p in self._grupo().items():
            datos[f'papel_{pk}'] = p['papel']
            datos[f'linea_{pk}'] = '' if p['linea'] is None else p['linea']
            for d in self.rotacion.dias_del_mes(2026, 11):
                datos[f'disp_{pk}_{d.isoformat()}'] = 'todos'
        respuesta = self.client.post(reverse('turnos:programar'), datos)
        self.assertRedirects(respuesta, f"{reverse('turnos:cuadro')}?anio=2026&mes=11")
        noviembre = AsignacionTurno.objects.filter(fecha__year=2026, fecha__month=11)
        self.assertFalse(noviembre.filter(funcionario=uno, tipo_turno__abreviatura='A').exists())
        primera_semana = noviembre.filter(funcionario=uno, fecha__day__lte=7)
        self.assertEqual({a.tipo_turno.abreviatura for a in primera_semana}, {'X1', 'T1'})  # su nueva línea

    def test_dos_titulares_con_la_misma_linea_no_se_programa(self):
        datos = {'cargo': self.cargo.pk, 'anio': 2026, 'mes': 11, 'accion': 'guardar'}
        for pk, p in self._grupo().items():
            datos[f'papel_{pk}'] = p['papel']
            datos[f'linea_{pk}'] = 0 if p['papel'] == 'titular' else ''
        respuesta = self.client.post(reverse('turnos:programar'), datos)
        self.assertContains(respuesta, 'Dos titulares tienen la misma línea')
        self.assertFalse(AsignacionTurno.objects.filter(fecha__month=11).exists())
