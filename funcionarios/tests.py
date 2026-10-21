import datetime
import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from .forms import FuncionarioForm
from .models import Area, Cargo, Funcionario

Usuario = get_user_model()
CLAVE = 'Clave-de-prueba-2026'


def crear_funcionario(documento, es_administrador=False, **extra):
    datos = {
        'tipo_documento': 'CC',
        'numero_documento': documento,
        'nombres': 'Ana',
        'apellidos': 'Pérez',
        'cargo': Cargo.objects.get_or_create(nombre='Auxiliar')[0].pk,
        'area': Area.objects.get_or_create(nombre='Urgencias')[0].pk,
        'es_administrador': es_administrador,
        'password1': CLAVE,
        'password2': CLAVE,
        **extra,
    }
    form = FuncionarioForm(datos)
    assert form.is_valid(), form.errors
    return form.save()


class RegistroFuncionarioTests(TestCase):
    def setUp(self):
        self.admin = crear_funcionario('1000', es_administrador=True)
        self.client.force_login(self.admin.usuario)
        self.cargo = Cargo.objects.get(nombre='Auxiliar')
        self.area = Area.objects.get(nombre='Urgencias')

    def _datos(self, **cambios):
        datos = {
            'tipo_documento': 'CC',
            'numero_documento': '2000',
            'nombres': 'Luis',
            'apellidos': 'Gómez',
            'cargo': self.cargo.pk,
            'area': self.area.pk,
            'correo': 'luis@example.com',
            'telefono': '3001234567',
            'fecha_ingreso': '2026-02-01',
            'password1': CLAVE,
            'password2': CLAVE,
        }
        datos.update(cambios)
        return datos

    def test_crear_funcionario_crea_su_usuario_y_puede_ingresar(self):
        respuesta = self.client.post(reverse('funcionarios:crear'), self._datos())
        self.assertRedirects(respuesta, reverse('funcionarios:lista'))

        funcionario = Funcionario.objects.get(numero_documento='2000')
        self.assertEqual(funcionario.usuario.username, '2000')
        self.assertEqual(funcionario.fecha_ingreso, datetime.date(2026, 2, 1))
        self.assertFalse(funcionario.es_administrador)

        self.client.logout()
        respuesta = self.client.post(reverse('ingresar'), {'username': '2000', 'password': CLAVE})
        self.assertRedirects(respuesta, reverse('calendario:mes'))

    def test_documento_repetido(self):
        respuesta = self.client.post(reverse('funcionarios:crear'), self._datos(numero_documento='1000'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.context['form'].errors['numero_documento'])

    def test_contrasenas_distintas(self):
        respuesta = self.client.post(reverse('funcionarios:crear'), self._datos(password2='otra-cosa-2026'))
        self.assertIn('password2', respuesta.context['form'].errors)
        self.assertFalse(Funcionario.objects.filter(numero_documento='2000').exists())

    def test_contrasena_obligatoria_al_crear(self):
        respuesta = self.client.post(reverse('funcionarios:crear'), self._datos(password1='', password2=''))
        self.assertIn('password1', respuesta.context['form'].errors)

    def test_editar_sin_contrasena_conserva_la_actual(self):
        otro = crear_funcionario('3000')
        respuesta = self.client.post(
            reverse('funcionarios:editar', args=[otro.pk]),
            self._datos(numero_documento='3001', password1='', password2=''),
        )
        self.assertRedirects(respuesta, reverse('funcionarios:lista'))
        otro.refresh_from_db()
        otro.usuario.refresh_from_db()
        self.assertEqual(otro.usuario.username, '3001')
        self.assertTrue(otro.usuario.check_password(CLAVE))

    def test_desactivar_impide_ingresar(self):
        otro = crear_funcionario('4000')
        self.client.post(reverse('funcionarios:estado', args=[otro.pk]))
        otro.usuario.refresh_from_db()
        self.assertFalse(otro.usuario.is_active)

        self.client.logout()
        respuesta = self.client.post(reverse('ingresar'), {'username': '4000', 'password': CLAVE})
        self.assertEqual(respuesta.status_code, 200)

    def test_no_puede_desactivarse_a_si_mismo(self):
        self.client.post(reverse('funcionarios:estado', args=[self.admin.pk]))
        self.admin.usuario.refresh_from_db()
        self.assertTrue(self.admin.usuario.is_active)

    def test_no_puede_quitarse_el_rol_de_administrador(self):
        respuesta = self.client.post(
            reverse('funcionarios:editar', args=[self.admin.pk]),
            self._datos(numero_documento='1000', password1='', password2=''),
        )
        self.assertIn('es_administrador', respuesta.context['form'].errors)

    def test_redireccion_externa_ignorada(self):
        otro = crear_funcionario('5000')
        respuesta = self.client.post(
            reverse('funcionarios:estado', args=[otro.pk]), {'siguiente': 'https://otro-sitio.com/'},
        )
        self.assertRedirects(respuesta, reverse('funcionarios:lista'))

    def test_filtros_del_listado(self):
        crear_funcionario('6000', nombres='Marta')
        inactivo = crear_funcionario('7000', nombres='Pedro')
        inactivo.usuario.is_active = False
        inactivo.usuario.save()

        activos = self.client.get(reverse('funcionarios:lista')).context['funcionarios']
        self.assertNotIn(inactivo, activos)
        inactivos = self.client.get(reverse('funcionarios:lista'), {'estado': 'inactivos'}).context['funcionarios']
        self.assertEqual(list(inactivos), [inactivo])
        busqueda = self.client.get(reverse('funcionarios:lista'), {'q': 'marta'}).context['funcionarios']
        self.assertEqual([f.numero_documento for f in busqueda], ['6000'])


class PermisosTests(TestCase):
    def test_funcionario_comun_no_entra_a_administracion(self):
        comun = crear_funcionario('8000')
        self.client.force_login(comun.usuario)
        self.assertEqual(self.client.get(reverse('funcionarios:lista')).status_code, 403)
        self.assertEqual(self.client.get(reverse('funcionarios:catalogo', args=['cargos'])).status_code, 403)
        self.assertEqual(self.client.get(reverse('calendario:mes')).status_code, 200)

    def test_sin_sesion_redirige_al_ingreso(self):
        respuesta = self.client.get(reverse('funcionarios:lista'))
        self.assertRedirects(respuesta, f"{reverse('ingresar')}?next={reverse('funcionarios:lista')}")


class CatalogoTests(TestCase):
    def setUp(self):
        self.client.force_login(crear_funcionario('1000', es_administrador=True).usuario)

    def test_agregar_y_nombre_repetido(self):
        url = reverse('funcionarios:catalogo', args=['areas'])
        self.client.post(url, {'nombre': '  Consulta   externa '})
        self.assertTrue(Area.objects.filter(nombre='Consulta externa').exists())
        respuesta = self.client.post(url, {'nombre': 'consulta externa'})
        self.assertTrue(respuesta.context['form'].errors)

    def test_cargo_inactivo_no_se_ofrece_en_formulario(self):
        Cargo.objects.create(nombre='Retirado', activo=False)
        form = FuncionarioForm()
        self.assertNotIn('Retirado', [c.nombre for c in form.fields['cargo'].queryset])

    def test_tipo_de_catalogo_desconocido(self):
        self.assertEqual(self.client.get('/funcionarios/configuracion/otros/').status_code, 404)


class ComandoFuncionariosPruebaTests(TestCase):
    def test_crear_y_borrar(self):
        import tempfile
        from django.core.management import call_command
        from django.test import override_settings
        from pathlib import Path

        Area.objects.create(nombre='Sala de Hemodialisis')
        with tempfile.TemporaryDirectory() as carpeta, override_settings(BASE_DIR=Path(carpeta)):
            call_command('funcionarios_prueba', stdout=io.StringIO())
            prueba = Funcionario.objects.filter(numero_documento__startswith='9000000')
            self.assertEqual(prueba.count(), 15)
            self.assertEqual(prueba.filter(cargo__nombre='Auxiliar de Enfermeria').count(), 7)
            self.assertEqual(prueba.filter(cargo__nombre='Enfermero Profesional').count(), 5)
            self.assertEqual(prueba.filter(cargo__nombre='Médico').count(), 3)
            self.assertTrue((Path(carpeta) / 'credenciales_prueba.txt').exists())

            call_command('funcionarios_prueba', stdout=io.StringIO())  # no duplica
            self.assertEqual(prueba.count(), 15)

            call_command('funcionarios_prueba', '--borrar', stdout=io.StringIO())
            self.assertEqual(prueba.count(), 0)
            self.assertFalse(Usuario.objects.filter(username__startswith='9000000').exists())


class ImportarExcelTests(TestCase):
    def setUp(self):
        self.admin = crear_funcionario('1000', es_administrador=True)
        self.client.force_login(self.admin.usuario)

    def _archivo(self, filas, nombre='funcionarios.xlsx'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        from openpyxl import load_workbook
        from .importar import HOJA_DATOS, plantilla_excel

        libro = load_workbook(io.BytesIO(plantilla_excel()))
        hoja = libro[HOJA_DATOS]
        for i, fila in enumerate(filas, start=2):
            for j, valor in enumerate(fila, start=1):
                hoja.cell(row=i, column=j, value=valor)
        salida = io.BytesIO()
        libro.save(salida)
        return SimpleUploadedFile(nombre, salida.getvalue())

    def _importar(self, filas):
        return self.client.post(reverse('funcionarios:importar'), {'archivo': self._archivo(filas)})

    def test_plantilla_se_descarga_con_las_listas(self):
        from openpyxl import load_workbook
        respuesta = self.client.get(reverse('funcionarios:plantilla'))
        self.assertEqual(respuesta.status_code, 200)
        libro = load_workbook(io.BytesIO(respuesta.content))
        self.assertEqual(libro.sheetnames, ['Instrucciones', 'Funcionarios', 'Listas'])
        self.assertEqual(libro['Funcionarios']['B1'].value, 'Número de documento *')
        self.assertEqual(libro['Listas']['B2'].value, 'Auxiliar')  # cargo existente

    def test_importa_y_crea_usuarios_con_contrasena_inicial(self):
        respuesta = self._importar([
            ['CC', 1012345678, 'Ana María', 'Pérez Gómez', 'Auxiliar', 'Urgencias', 'ana@correo.com', 3001234567,
             datetime.date(2024, 2, 1), 'No'],
            ['Cédula de ciudadanía', '2030', 'Luis', 'Mora', 'Jefe de turno', 'Urgencias', '', '', '01/03/2025', 'Sí'],
        ])
        self.assertEqual(respuesta.status_code, 200)
        credenciales = respuesta.context['credenciales']
        self.assertEqual([c[0] for c in credenciales], ['1012345678', '2030'])
        ana = Funcionario.objects.get(numero_documento='1012345678')
        self.assertEqual((ana.nombres, ana.telefono, ana.fecha_ingreso, ana.es_administrador),
                         ('Ana María', '3001234567', datetime.date(2024, 2, 1), False))
        self.assertTrue(ana.usuario.check_password(credenciales[0][2]))
        luis = Funcionario.objects.get(numero_documento='2030')
        self.assertTrue(luis.es_administrador)
        self.assertEqual(luis.cargo.nombre, 'Jefe de turno')  # cargo nuevo creado
        self.assertIn('Cargo: Jefe de turno', respuesta.context['nuevos'])

        # Las contraseñas se descargan una sola vez.
        descarga = self.client.get(reverse('funcionarios:credenciales'))
        self.assertEqual(descarga['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        self.assertRedirects(self.client.get(reverse('funcionarios:credenciales')), reverse('funcionarios:lista'))

    def test_si_una_fila_falla_no_se_crea_ninguno(self):
        antes = Funcionario.objects.count()
        respuesta = self._importar([
            ['CC', '3001', 'Bien', 'Escrito', 'Auxiliar', 'Urgencias', '', '', '', ''],
            ['CC', '3002', '', 'Sin nombre', 'Nuevo cargo', 'Urgencias', 'correo-malo', '', '', ''],
            ['CC', '3001', 'Repetido', 'En archivo', 'Auxiliar', 'Urgencias', '', '', '', ''],
            ['CC', '1000', 'Ya', 'Existe', 'Auxiliar', 'Urgencias', '', '', '', ''],
        ])
        errores = respuesta.context['errores']
        self.assertEqual({fila for fila, _, _ in errores}, {3, 4, 5})
        self.assertEqual(Funcionario.objects.count(), antes)
        self.assertFalse(Cargo.objects.filter(nombre='Nuevo cargo').exists())
        self.assertContains(respuesta, 'No se creó ningún funcionario')

    def test_archivo_invalido(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        respuesta = self.client.post(reverse('funcionarios:importar'), {'archivo': SimpleUploadedFile('x.xlsx', b'no es excel')})
        self.assertContains(respuesta, 'No se pudo leer el archivo')
        respuesta = self.client.post(reverse('funcionarios:importar'), {'archivo': SimpleUploadedFile('x.csv', b'a,b')})
        self.assertContains(respuesta, 'debe ser de Excel')
        self.assertContains(self._importar([]), 'no tiene funcionarios')

    def test_solo_administrador(self):
        comun = crear_funcionario('5000')
        self.client.force_login(comun.usuario)
        self.assertEqual(self.client.get(reverse('funcionarios:importar')).status_code, 403)
        self.assertEqual(self.client.get(reverse('funcionarios:plantilla')).status_code, 403)
