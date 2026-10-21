"""Plantilla de Excel para cargar funcionarios reales y su importación al sistema."""
import datetime
import io
import secrets

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .forms import FuncionarioForm
from .models import Area, Cargo, Funcionario

HOJA_DATOS = 'Funcionarios'
MAX_FILAS = 500

# (clave, encabezado, obligatorio, ancho, ayuda)
COLUMNAS = [
    ('tipo_documento', 'Tipo de documento', True, 18, 'CC, CE, TI, PA o PPT. Elija de la lista.'),
    ('numero_documento', 'Número de documento', True, 20, 'Solo el número, sin puntos. También será su usuario para entrar.'),
    ('nombres', 'Nombres', True, 22, ''),
    ('apellidos', 'Apellidos', True, 24, ''),
    ('cargo', 'Cargo', True, 26, 'Elija de la lista. Si escribe uno nuevo, se crea al importar.'),
    ('area', 'Área', True, 24, 'Elija de la lista. Si escribe una nueva, se crea al importar.'),
    ('correo', 'Correo', False, 28, ''),
    ('telefono', 'Teléfono', False, 16, ''),
    ('fecha_ingreso', 'Fecha de ingreso', False, 18, 'Formato fecha, por ejemplo 01/02/2024.'),
    ('es_administrador', 'Es administrador', False, 18, 'Sí o No. Si se deja vacío, No.'),
]

EJEMPLO = ['CC', '1012345678', 'Ana María', 'Pérez Gómez', 'Auxiliar de Enfermeria', 'Sala de Hemodialisis',
           'ana.perez@correo.com', '3001234567', datetime.date(2024, 2, 1), 'No']

FUENTE = 'Arial'
AZUL = '0770B8'
AZUL_MARINO = '1F2A5C'
AMARILLO = 'FFF8D6'


def _fuente(**kwargs):
    return Font(name=FUENTE, **kwargs)


def plantilla_excel():
    """Devuelve el archivo de plantilla (bytes) con las listas de cargos y áreas actuales."""
    libro = Workbook()
    borde = Border(bottom=Side(style='thin', color='D6E4EC'))

    # --- Hoja de datos (la que se llena) ---
    hoja = libro.active
    hoja.title = HOJA_DATOS
    for col, (_, encabezado, obligatorio, ancho, ayuda) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=col, value=encabezado + (' *' if obligatorio else ''))
        celda.font = _fuente(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL)
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if ayuda:
            celda.comment = Comment(ayuda, 'Control de Horas')
        hoja.column_dimensions[get_column_letter(col)].width = ancho
        # Celdas para llenar: fondo amarillo suave.
        for fila in range(2, MAX_FILAS + 2):
            dato = hoja.cell(row=fila, column=col)
            dato.fill = PatternFill('solid', fgColor=AMARILLO)
            dato.font = _fuente()
            dato.border = borde
    hoja.row_dimensions[1].height = 30
    hoja.freeze_panes = 'A2'
    for fila in range(2, MAX_FILAS + 2):
        hoja.cell(row=fila, column=2).number_format = '@'  # documento como texto (no pierde ceros)
        hoja.cell(row=fila, column=8).number_format = '@'  # teléfono como texto
        hoja.cell(row=fila, column=9).number_format = 'DD/MM/YYYY'

    # --- Hoja de listas (para las listas desplegables) ---
    listas = libro.create_sheet('Listas')
    tipos = [codigo for codigo, _ in Funcionario.TipoDocumento.choices]
    cargos = list(Cargo.objects.filter(activo=True).values_list('nombre', flat=True))
    areas = list(Area.objects.filter(activo=True).values_list('nombre', flat=True))
    for col, (titulo, valores) in enumerate(
        [('Tipo de documento', tipos), ('Cargos', cargos), ('Áreas', areas), ('Sí/No', ['Sí', 'No'])], start=1,
    ):
        cabecera = listas.cell(row=1, column=col, value=titulo)
        cabecera.font = _fuente(bold=True)
        for fila, valor in enumerate(valores, start=2):
            listas.cell(row=fila, column=col, value=valor).font = _fuente()
        listas.column_dimensions[get_column_letter(col)].width = 28

    rango = f'2:{MAX_FILAS + 1}'
    validaciones = [
        ('A', f'=Listas!$A$2:$A${len(tipos) + 1}', False),
        ('E', f'=Listas!$B$2:$B${max(len(cargos), 1) + 1}', True),
        ('F', f'=Listas!$C$2:$C${max(len(areas), 1) + 1}', True),
        ('J', '=Listas!$D$2:$D$3', False),
    ]
    for columna, formula, permite_otros in validaciones:
        validacion = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=not permite_otros)
        validacion.error = 'Elija un valor de la lista.'
        hoja.add_data_validation(validacion)
        validacion.add(f'{columna}{rango.split(":")[0]}:{columna}{rango.split(":")[1]}')

    # --- Hoja de instrucciones ---
    instrucciones = libro.create_sheet('Instrucciones', 0)
    instrucciones.column_dimensions['A'].width = 26
    instrucciones.column_dimensions['B'].width = 90
    titulo = instrucciones.cell(row=1, column=1, value='Plantilla para cargar funcionarios')
    titulo.font = _fuente(bold=True, size=16, color=AZUL_MARINO)
    pasos = [
        ('1.', f'Llene una fila por funcionario en la hoja «{HOJA_DATOS}» (celdas amarillas). No cambie los encabezados.'),
        ('2.', 'Las columnas con * son obligatorias. Pase el mouse sobre cada encabezado para ver una ayuda.'),
        ('3.', 'Tipo de documento, Cargo, Área y Es administrador tienen lista desplegable.'),
        ('4.', 'No escriba contraseñas: el sistema genera una contraseña inicial para cada funcionario al importar.'),
        ('5.', 'En el sistema: Funcionarios → Importar desde Excel → elija este archivo.'),
        ('6.', 'Si alguna fila tiene un error, no se crea ningún funcionario y el sistema le dice qué corregir.'),
        ('', ''),
        ('Importante', 'Este archivo tiene datos personales: guárdelo fuera de la carpeta del proyecto y bórrelo al terminar.'),
    ]
    for fila, (a, b) in enumerate(pasos, start=3):
        instrucciones.cell(row=fila, column=1, value=a).font = _fuente(bold=True)
        celda = instrucciones.cell(row=fila, column=2, value=b)
        celda.font = _fuente(bold=(a == 'Importante'), color='C0392B' if a == 'Importante' else None)
        celda.alignment = Alignment(wrap_text=True, vertical='top')

    fila_ejemplo = len(pasos) + 5
    instrucciones.cell(row=fila_ejemplo, column=1, value='Ejemplo de una fila').font = _fuente(bold=True, color=AZUL_MARINO)
    for i, ((_, encabezado, *_), valor) in enumerate(zip(COLUMNAS, EJEMPLO)):
        instrucciones.cell(row=fila_ejemplo + 1 + i, column=1, value=encabezado).font = _fuente(bold=True)
        celda = instrucciones.cell(row=fila_ejemplo + 1 + i, column=2, value=valor)
        celda.font = _fuente(italic=True)
        celda.alignment = Alignment(horizontal='left')
        if isinstance(valor, datetime.date):
            celda.number_format = 'DD/MM/YYYY'

    libro.active = 1  # abre en la hoja de datos
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)  # documentos/teléfonos que Excel guardó como número
    return ' '.join(str(valor).split())


def _fecha(valor):
    if valor in (None, ''):
        return ''
    if isinstance(valor, datetime.datetime):
        return valor.date().isoformat()
    if isinstance(valor, datetime.date):
        return valor.isoformat()
    texto = _texto(valor)
    for formato in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return texto  # el formulario dirá que no es una fecha válida


def _si_no(valor):
    return _texto(valor).lower() in ('sí', 'si', 's', 'x', 'yes', '1', 'true', 'verdadero')


def _tipo_documento(valor):
    texto = _texto(valor)
    for codigo, nombre in Funcionario.TipoDocumento.choices:
        if texto.upper() == codigo or texto.lower() == nombre.lower():
            return codigo
    return texto


def contrasena_inicial():
    """Contraseña inicial legible y aleatoria (sin caracteres que se confundan)."""
    letras = 'abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ'
    numeros = '23456789'
    return (''.join(secrets.choice(letras) for _ in range(5)) + '-' + ''.join(secrets.choice(numeros) for _ in range(4)))


class ErrorDeArchivo(Exception):
    """El archivo no se puede leer como plantilla."""


def leer_filas(archivo):
    """Lee la hoja de funcionarios y devuelve [(número de fila, {campo: valor}), …] sin las filas vacías."""
    try:
        libro = load_workbook(archivo, data_only=True, read_only=True)
    except Exception:
        raise ErrorDeArchivo('No se pudo leer el archivo. Use la plantilla en formato .xlsx.')
    if HOJA_DATOS not in libro.sheetnames:
        raise ErrorDeArchivo(f'El archivo no tiene la hoja «{HOJA_DATOS}». Use la plantilla del sistema.')
    hoja = libro[HOJA_DATOS]
    filas = []
    for numero, valores in enumerate(hoja.iter_rows(min_row=2, max_col=len(COLUMNAS), values_only=True), start=2):
        if all(_texto(v) == '' for v in valores):
            continue
        filas.append((numero, dict(zip([c[0] for c in COLUMNAS], valores))))
    libro.close()
    if not filas:
        raise ErrorDeArchivo(f'La hoja «{HOJA_DATOS}» no tiene funcionarios.')
    if len(filas) > MAX_FILAS:
        raise ErrorDeArchivo(f'El archivo tiene {len(filas)} funcionarios; el máximo por archivo es {MAX_FILAS}.')
    return filas


def importar(archivo):
    """Crea los funcionarios del archivo. Todo o nada: si una fila tiene errores, no se crea ninguno.

    Devuelve (creados, errores, nuevos_catalogos):
    - creados: [(funcionario, contraseña inicial), …]
    - errores: [(fila, campo, mensaje), …]
    - nuevos_catalogos: nombres de cargos y áreas que se crearon.
    """
    filas = leer_filas(archivo)
    nombres_campos = {clave: encabezado for clave, encabezado, *_ in COLUMNAS}
    creados, errores, nuevos = [], [], []
    documentos_vistos = {}

    with transaction.atomic():
        for numero, datos in filas:
            documento = _texto(datos['numero_documento'])
            if documento and documento in documentos_vistos:
                errores.append((numero, 'Número de documento', f'Repetido en el archivo (también en la fila {documentos_vistos[documento]}).'))
                continue
            documentos_vistos[documento] = numero

            catalogos = {}
            catalogo_inactivo = False
            for campo, modelo in (('cargo', Cargo), ('area', Area)):
                nombre = _texto(datos[campo])
                if not nombre:
                    catalogos[campo] = ''
                    continue
                objeto = modelo.objects.filter(nombre__iexact=nombre).first()
                if objeto is None:
                    objeto = modelo.objects.create(nombre=nombre)
                    nuevos.append(f'{modelo._meta.verbose_name.capitalize()}: {nombre}')
                elif not objeto.activo:
                    errores.append((numero, nombres_campos[campo], f'«{objeto.nombre}» está desactivado en el sistema.'))
                    catalogo_inactivo = True
                catalogos[campo] = objeto.pk
            if catalogo_inactivo:
                continue

            clave = contrasena_inicial()
            form = FuncionarioForm({
                'tipo_documento': _tipo_documento(datos['tipo_documento']),
                'numero_documento': documento,
                'nombres': _texto(datos['nombres']),
                'apellidos': _texto(datos['apellidos']),
                'cargo': catalogos['cargo'],
                'area': catalogos['area'],
                'correo': _texto(datos['correo']),
                'telefono': _texto(datos['telefono']),
                'fecha_ingreso': _fecha(datos['fecha_ingreso']),
                'es_administrador': _si_no(datos['es_administrador']),
                'password1': clave,
                'password2': clave,
            })
            if form.is_valid():
                creados.append((form.save(), clave))
            else:
                for campo, mensajes in form.errors.items():
                    for mensaje in mensajes:
                        errores.append((numero, nombres_campos.get(campo, 'Fila'), mensaje))

        if errores:
            transaction.set_rollback(True)
            return [], errores, []
    return creados, errores, nuevos


def credenciales_excel(credenciales):
    """Excel con documento, nombre y contraseña inicial de los funcionarios recién importados."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Contraseñas iniciales'
    encabezados = ['Número de documento (usuario)', 'Nombre', 'Contraseña inicial']
    for col, texto in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=col, value=texto)
        celda.font = _fuente(bold=True, color='FFFFFF')
        celda.fill = PatternFill('solid', fgColor=AZUL)
        hoja.column_dimensions[get_column_letter(col)].width = 32
    for fila, (documento, nombre, clave) in enumerate(credenciales, start=2):
        for col, valor in enumerate((documento, nombre, clave), start=1):
            hoja.cell(row=fila, column=col, value=valor).font = _fuente()
    nota = hoja.cell(row=len(credenciales) + 3, column=1,
                     value='Entregue cada contraseña solo a su dueño y borre este archivo después.')
    nota.font = _fuente(italic=True, color='C0392B')
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
